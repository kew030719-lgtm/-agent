from itertools import zip_longest
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from python_calamine import CalamineWorkbook

from app.knowledge.blocks import ParsedBlock, SourceFormat, SourceLocation


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _cached_or_formula(cached: object, formula: object) -> object:
    if cached is not None:
        return cached
    if isinstance(formula, str) and formula.startswith("="):
        return f"公式（无缓存值）：{formula}"
    return formula


def _sheet_blocks(
    sheet_name: str,
    rows: Iterable[tuple[int, tuple[object, ...]]],
    source_format: SourceFormat,
    start_order: int,
) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    headers: tuple[str, ...] = ()
    table_index = 0
    for row_number, raw_values in rows:
        values = tuple(_cell_text(value) for value in raw_values)
        if not any(values):
            headers = ()
            continue
        if not headers:
            table_index += 1
            headers = tuple(
                value or f"列{column}"
                for column, value in enumerate(values, start=1)
            )
            blocks.append(ParsedBlock(
                source_format=source_format,
                order=start_order + len(blocks), kind="paragraph",
                text=f"工作表 {sheet_name}，表格 {table_index} 列名："
                     + "、".join(headers),
                location=SourceLocation(
                    sheet_name=sheet_name, table_index=table_index,
                    row_start=row_number, row_end=row_number,
                ),
                table_headers=headers, text_source="native",
            ))
            continue
        content = "；".join(
            f"{headers[column] if column < len(headers) else f'列{column + 1}'}：{value}"
            for column, value in enumerate(values) if value
        )
        blocks.append(ParsedBlock(
            source_format=source_format,
            order=start_order + len(blocks), kind="table_row",
            text=content,
            location=SourceLocation(
                sheet_name=sheet_name, table_index=table_index,
                row_start=row_number, row_end=row_number,
            ),
            table_headers=headers, text_source="native",
        ))
    return blocks


def parse_excel(path: Path) -> list[ParsedBlock]:
    suffix = path.suffix.lower()
    blocks: list[ParsedBlock] = []
    if suffix == ".xls":
        with CalamineWorkbook.from_path(str(path)) as workbook:
            for sheet_name in workbook.sheet_names:
                sheet = workbook.get_sheet_by_name(sheet_name)
                rows = (
                    (number, tuple(values))
                    for number, values in enumerate(
                        sheet.to_python(skip_empty_area=False), start=1
                    )
                )
                blocks.extend(_sheet_blocks(
                    sheet_name, rows, "xls", len(blocks)
                ))
        return blocks

    cached_book = load_workbook(path, read_only=True, data_only=True)
    formula_book = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet_name in cached_book.sheetnames:
            cached_rows = cached_book[sheet_name].iter_rows(values_only=True)
            formula_rows = formula_book[sheet_name].iter_rows(values_only=True)
            rows = (
                (
                    number,
                    tuple(
                        _cached_or_formula(cached, formula)
                        for cached, formula in zip_longest(
                            cached_row or (), formula_row or ()
                        )
                    ),
                )
                for number, (cached_row, formula_row) in enumerate(
                    zip_longest(cached_rows, formula_rows), start=1
                )
            )
            blocks.extend(_sheet_blocks(
                sheet_name, rows, suffix[1:], len(blocks)
            ))
    finally:
        cached_book.close()
        formula_book.close()
    return blocks
