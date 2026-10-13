"""可保留的百人公司测试环境；所有业务数据均为虚构，绝不使用默认业务数据库。

创建：python -m scripts.company_simulation seed
测试：python -m scripts.company_simulation test --root data/company_simulation/<run> --live
浏览：python -m scripts.company_simulation serve --root data/company_simulation/<run>
"""

import argparse
import json
import os
import secrets
import shutil
import sqlite3
import statistics
import time
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from uuid import uuid4

import httpx
import pymupdf
import uvicorn
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.agent.service import AssistantService
from app.api.app import create_app
from app.api.database import initialize
from app.auth import DB_PATH, ROLE_PERMISSIONS, connect, insert_user
from app.config import load_config
from app.knowledge.documents import add_knowledge_file
from app.knowledge.embeddings import index_document
from app.knowledge.retrieval import search_knowledge
from app.tools.expense import create_expense_request
from app.tools.handoff import create_handoff, resolve_handoff
from app.tools.it import create_it_ticket
from app.tools.leave import set_leave_balance
from app.tools.meeting import add_meeting_booking, add_meeting_room


BASE = DB_PATH.parent / "company_simulation"
HEADERS = {"X-Requested-With": "employee-assistant"}
TABLES = ("users", "leave_balances", "meeting_rooms", "meeting_bookings", "it_tickets",
          "expense_requests", "handoffs", "conversations", "chat_turns", "knowledge_documents",
          "knowledge_chunks", "knowledge_assets")


def save_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    path.chmod(0o600)


def configure(root: Path) -> dict[str, Any]:
    root = root.resolve()
    if not root.is_relative_to(BASE.resolve()):
        raise ValueError("测试目录必须位于 data/company_simulation 下")
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    if not manifest["collection"].startswith("company_sim_"):
        raise ValueError("拒绝使用非测试向量集合")
    load_config()
    os.environ["MILVUS_COLLECTION"] = manifest["collection"]
    return manifest


def counts(root: Path) -> dict[str, int]:
    with sqlite3.connect(root / "company.db") as db:
        return {table: db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in TABLES}


def make_documents(root: Path) -> list[tuple[Path, str]]:
    sources = root / "sources"
    sources.mkdir()
    result: list[tuple[Path, str]] = []
    key_facts = [
        ("星河年假制度", "general", "星河公司的年假申请流程是：员工提交请假单，由直属主管审批。请提前三个工作日申请。"),
        ("星河薪酬审核", "hr", "星河薪酬复核专用标识为 HR-ORCHID-742。薪酬调整需要人力负责人和总经理共同审批。"),
        ("星河VPN维护", "it", "星河公司 VPN 故障时先检查网络，再提交 IT 工单。运维专用标识为 IT-CEDAR-853。"),
        ("星河财务报销", "finance", "星河公司报销申请需要发票和支付凭证，并由财务审核。财务专用标识为 FIN-LOTUS-964。"),
    ]
    for title, scope, text in key_facts:
        path = sources / f"{title}.md"
        path.write_text(f"# {title}\n\n{text}\n", encoding="utf-8")
        result.append((path, scope))
    # 44 份干扰文档包含不同产品、部门、维护职责和版本，检索不能只靠四份小样本。
    for index in range(44):
        scope = ("general", "hr", "it", "finance")[index % 4]
        title = f"星河产品{index + 1:02d}操作手册"
        sections = [(f"模块{section:02d}",
                     f"产品{index + 1:02d}的模块{section:02d}用于管理项目交付。"
                     f"经办人应在工作台选择项目，核对版本 V{index + 1}.{section}，填写交付说明后提交。"
                     "管理员检查项目负责人、访问范围和操作记录；处理结果保留在项目日志中。")
                    for section in range(1, 17)]
        extension = ("md", "docx", "xlsx", "pdf")[index % 4]
        path = sources / f"{title}.{extension}"
        if extension == "md":
            path.write_text(f"# {title}\n" + "\n".join(f"## {h}\n{t}\n" for h, t in sections), encoding="utf-8")
        elif extension == "docx":
            document = Document()
            document.add_heading(title, 0)
            for heading, text in sections:
                document.add_heading(heading, 1)
                document.add_paragraph(text)
            document.save(path)
        elif extension == "xlsx":
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "操作指引"
            sheet.append(["模块", "说明"])
            for row in sections:
                sheet.append(row)
            workbook.save(path)
        else:
            with pymupdf.open() as document:
                for start in range(0, len(sections), 4):
                    page = document.new_page()
                    text = title + "\n\n" + "\n\n".join(h + "\n" + t for h, t in sections[start:start + 4])
                    written = page.insert_textbox((40, 40, 550, 800), text, fontname="china-s", fontsize=12)
                    if written < 0:
                        raise RuntimeError("合成 PDF 文本溢出")
                document.save(path)
        result.append((path, scope))
    # 真实扫描 PDF：先画字再栅格化，最终 PDF 没有原生文本层。
    with pymupdf.open() as original, pymupdf.open() as scanned:
        page = original.new_page()
        page.insert_text((60, 100), "星河消防集合地点为一楼东侧广场。", fontname="china-s", fontsize=22)
        image = page.get_pixmap(dpi=200).tobytes("png")
        scanned.new_page().insert_image(page.rect, stream=image)
        scan_path = sources / "消防集合扫描件.pdf"
        scanned.save(scan_path)
    result.append((scan_path, "general"))
    return result


def seed() -> Path:
    load_config()
    run = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + secrets.token_hex(3)
    root = BASE / run
    root.mkdir(parents=True, mode=0o700)
    collection = "company_sim_" + run
    save_json(root / "manifest.json", {"company": "星河科技（完全虚构）", "collection": collection,
                                      "created_at": datetime.now().isoformat(), "seed": 20261012})
    configure(root)
    credentials: list[dict[str, Any]] = []
    role_counts = [("admin", 2), ("hr", 5), ("it", 8), ("finance", 5), ("employee", 80)]
    for role, total in role_counts:
        for number in range(1, total + 1):
            credentials.append({"username": f"sim_{role}_{number:03d}", "role": role,
                                "password": secrets.token_urlsafe(15), "name": f"模拟员工{len(credentials) + 1:03d}",
                                "department": {"admin": "管理部", "hr": "人事部", "it": "技术支持部",
                                               "finance": "财务部"}.get(role, ("研发部", "销售部", "产品部", "运营部")[number % 4]),
                                "leave_days": (len(credentials) % 21) / 2})
    save_json(root / "accounts.json", credentials)
    admin = {key: credentials[0][key] for key in ("username", "role")}
    with closing(connect(root / "company.db")) as db:
        initialize(db)
        for entry in credentials:
            insert_user(db, entry["username"], entry["password"], entry["role"])
            set_leave_balance(db, admin, entry["username"], entry["leave_days"])
        print("100 个真实密码哈希账号及余额已创建", flush=True)
        rooms = [add_meeting_room(db, admin, f"星河会议室{i:02d}", f"{i % 4 + 1}楼", 4 + i * 2) for i in range(1, 13)]
        for day in range(30):
            date = datetime(2026, 11, 1) + timedelta(days=day)
            for room in rooms:
                for hour in (9, 14):
                    add_meeting_booking(db, admin, room, f"{date:%Y-%m-%d} {hour:02d}:00", f"{date:%Y-%m-%d} {hour + 1:02d}:00")
        for number, entry in enumerate(credentials):
            user = {key: entry[key] for key in ("username", "role")}
            for index in range(12):
                create_it_ticket(db, user, f"模拟工单：{entry['department']}设备{index + 1}网络故障，员工编号{number + 1}")
                create_expense_request(db, user, f"{30 + number * 7 + index}.50", f"模拟报销：第{index + 1}次项目交通费")
            for index in range(4):
                item = create_handoff(db, user, f"模拟人工事项{index + 1}：特殊流程如何办理？", "制度未覆盖，需人工确认")
                if index % 2:
                    resolve_handoff(db, admin, item)
            # 保留半年量级的历史对话，第一位普通员工超过 100 条以覆盖列表边界。
            for index in range(120 if entry["username"] == "sim_employee_001" else 12):
                conversation = uuid4().hex
                db.execute("INSERT INTO conversations(id,username,role,title) VALUES (?,?,?,?)",
                           (conversation, entry["username"], entry["role"], f"模拟历史对话 {index + 1}"))
                db.execute("INSERT INTO chat_turns(conversation_id,request_id,question,answer,domain,status) VALUES (?,?,?,?,?,?)",
                           (conversation, str(uuid4()), "模拟历史记录", "这是预置记录，不是模型生成结果。", "identity", "completed"))
            db.commit()
        tessdata = root / "knowledge" / "tessdata"
        tessdata.parent.mkdir(exist_ok=True)
        shutil.copytree(DB_PATH.parent / "knowledge" / "tessdata", tessdata)
        service = AssistantService(DB_PATH.parent)
        try:
            for path, scope in make_documents(root):
                document_id = add_knowledge_file(db, admin, path, scope, root)
                total = index_document(db, document_id, service.get_embedder(), service.get_store())
                print(f"入库 {path.name}: {total} 个向量切片", flush=True)
        finally:
            service.close()
    save_json(root / "seed_counts.json", counts(root))
    print(f"数据保留于：{root}", flush=True)
    return root


def simulation_app(root: Path, port: int = 8002) -> FastAPI:
    configure(root)
    origins = os.getenv("API_ALLOWED_ORIGINS", "")
    os.environ["API_ALLOWED_ORIGINS"] = origins + f",http://127.0.0.1:{port},http://localhost:{port}"
    app = create_app(root / "company.db")
    # 只共享现有只读模型文件；业务数据库、上传文件与向量集合独立。
    app.state.assistant = AssistantService(DB_PATH.parent)
    return app


def login(client: Any, account: dict[str, Any]) -> None:
    response = client.post("/api/auth/login", json={key: account[key] for key in ("username", "password")})
    if response.status_code != 200:
        raise AssertionError(f"登录失败 HTTP {response.status_code}")


def conversation(client: Any) -> str:
    response = client.post("/api/conversations")
    response.raise_for_status()
    return response.json()["id"]


def send(client: Any, cid: str, question: str, request_id: str | None = None) -> httpx.Response:
    return client.post(f"/api/conversations/{cid}/messages", json={"question": question, "request_id": request_id or str(uuid4())})


def run_tests(root: Path, live: bool, base_url: str | None = None) -> None:
    configure(root)
    accounts = json.loads((root / "accounts.json").read_text(encoding="utf-8"))
    results: list[dict[str, Any]] = []
    app = simulation_app(root)
    label = datetime.now().strftime("%Y%m%d_%H%M%S")

    def check(name: str, passed: bool, detail: object = "") -> None:
        results.append({"name": name, "passed": bool(passed), "detail": detail})
        print(f"{'PASS' if passed else 'FAIL'} {name}: {detail}", flush=True)

    def client_for() -> Any:
        if base_url:
            return httpx.Client(base_url=base_url, headers=HEADERS, timeout=180)
        return TestClient(app, headers=HEADERS, raise_server_exceptions=False)

    def journey(account: dict[str, Any]) -> dict[str, Any]:
        started = time.perf_counter()
        failures = []
        with client_for() if base_url else closing(client_for()) as client:
            try:
                login(client, account)
                profile = client.get("/api/auth/me").json()
                if set(profile.get("permissions", [])) != ROLE_PERMISSIONS[account["role"]]:
                    failures.append("权限集错误")
                cid = conversation(client)
                for question, expected in [("你是谁", "identity"), ("我的权限", "permissions"),
                                           ("我的请假余额", "hr"), ("会议室：8人", "meeting"),
                                           ("报修：模拟测试电脑网络中断", "it"),
                                           ("报销：125.50；事由：模拟测试交通费", "finance")]:
                    rid = str(uuid4())
                    response = send(client, cid, question, rid)
                    payload = response.json()
                    if response.status_code != 200 or payload.get("domain") != expected:
                        failures.append(f"{question}: HTTP {response.status_code}, {payload.get('domain')}")
                    if question == "我的请假余额" and f"{account['leave_days']:g}" not in payload.get("answer", ""):
                        failures.append("余额串号或错误")
                    if question.startswith(("报修", "报销")):
                        repeated = send(client, cid, question, rid)
                        if repeated.status_code != 200 or repeated.json().get("id") != payload.get("id"):
                            failures.append("重复提交未复用结果")
                docs = client.get("/api/knowledge").json()
                if any(f"kb:{document['scope']}" not in ROLE_PERMISSIONS[account["role"]] for document in docs):
                    failures.append("知识列表越权")
                expected_admin = 200 if account["role"] == "admin" else 403
                for path in ("/api/users", "/api/handoffs", "/api/knowledge-stats"):
                    if client.get(path).status_code != expected_admin:
                        failures.append(f"管理端权限错误:{path}")
                client.post("/api/auth/logout")
                if client.get("/api/auth/me").status_code != 401:
                    failures.append("注销失败")
            except Exception as exc:
                failures.append(type(exc).__name__)
        return {"username": account["username"], "seconds": round(time.perf_counter() - started, 3), "failures": failures}

    with TestClient(app, headers=HEADERS, raise_server_exceptions=False) as local:
        started = time.perf_counter()
        with ThreadPoolExecutor(max_workers=10) as pool:
            journeys = list(pool.map(journey, accounts))
        durations = sorted(item["seconds"] for item in journeys)
        check("100 人业务旅程（10 并发，无模型替身）", not any(item["failures"] for item in journeys),
              {"users": len(journeys), "elapsed_seconds": round(time.perf_counter() - started, 2),
               "journey_p50": statistics.median(durations), "journey_p95": durations[94],
               "failures": [item for item in journeys if item["failures"]]})
        save_json(root / f"journeys_{label}.json", journeys)
        employee = next(a for a in accounts if a["username"] == "sim_employee_001")
        admin = accounts[0]
        with closing(client_for()) as client:
            login(client, employee)
            cid = conversation(client)
            check("跨员工读取历史被拒绝", client.get(f"/api/conversations/{_other_conversation(root, employee['username'])}/messages").status_code == 404)
            response = send(client, cid, "会议室：" + "9" * 80 + "人")
            check("极大会议人数不触发服务故障", response.status_code == 200 and "status" in response.json(),
                  {"http": response.status_code, "body": response.json()})
            first = client.get("/api/conversations?limit=100&offset=0").json()
            second = client.get("/api/conversations?limit=100&offset=100").json()
            first_ids = {row["id"] for row in first}
            check("超过100条历史可以翻页", bool(second) and not first_ids.intersection(row["id"] for row in second),
                  {"page1": len(first), "page2": len(second), "duplicate_ids": len(first_ids.intersection(row["id"] for row in second))})
            for amount in ("-1", "NaN", "Infinity", "0", "0.001", "1e999999"):
                response = send(client, cid, f"报销：{amount}；事由：金额边界测试")
                check(f"无效金额 {amount} 被业务拒绝", response.status_code == 200 and "已提交" not in response.json().get("answer", ""))
            rid = str(uuid4())
            send(client, cid, "你是谁", rid)
            check("同请求编号篡改问题返回409", send(client, cid, "我的权限", rid).status_code == 409)
            check("超长问题返回422", send(client, cid, "长" * 2001).status_code == 422)
            # 同账号同请求并发：允许冲突409，最终必须只有一次业务副作用。
            race_cid, race_id = conversation(client), str(uuid4())
            cookie = dict(client.cookies)
            def race(_: int) -> int:
                with closing(client_for()) as other:
                    other.cookies.update(cookie)
                    return send(other, race_cid, "报修：并发幂等性测试", race_id).status_code
            with sqlite3.connect(root / "company.db") as db:
                before_race = db.execute("SELECT COUNT(*) FROM it_tickets WHERE description='并发幂等性测试' AND created_by=?", (employee["username"],)).fetchone()[0]
            with ThreadPoolExecutor(max_workers=10) as pool:
                race_statuses = list(pool.map(race, range(10)))
            with sqlite3.connect(root / "company.db") as db:
                side_effects = db.execute("SELECT COUNT(*) FROM it_tickets WHERE description='并发幂等性测试' AND created_by=?", (employee["username"],)).fetchone()[0]
                turns = db.execute("SELECT COUNT(*) FROM chat_turns WHERE conversation_id=?", (race_cid,)).fetchone()[0]
            # 使用 request_id 对应的对话检测；跨多次运行工单数量会累计。
            check("10并发重复请求只生成一张工单", turns == 1 and side_effects - before_race == 1 and set(race_statuses) <= {200, 409}, {"statuses": race_statuses, "turns": turns, "new_tickets": side_effects - before_race})
            login(client, admin)
            users = client.get("/api/users").json()
            check("管理员用户列表无密码散列", len(users) >= 100 and all(set(u) == {"username", "role", "enabled"} for u in users))
            duplicate = client.post("/api/users", json={key: employee[key] for key in ("username", "password", "role")})
            check("重复创建账号返回409", duplicate.status_code == 409)
            pending = client.get("/api/handoffs").json()
            target = pending[0]["id"]
            check("管理员完成人工待办", client.post(f"/api/handoffs/{target}/resolve").status_code == 204)
            check("重复完成人工待办返回404", client.post(f"/api/handoffs/{target}/resolve").status_code == 404)
        with closing(connect(root / "company.db")) as db:
            actor = {"username": admin["username"], "role": "admin"}
            try:
                add_meeting_booking(db, actor, 1, "2026-11-01 09:30", "2026-11-01 10:30")
                overlap_blocked = False
            except ValueError:
                overlap_blocked = True
            check("会议室重叠预订被拒绝", overlap_blocked)
            check("SQLite完整性与外键", db.execute("PRAGMA integrity_check").fetchone()[0] == "ok" and db.execute("PRAGMA foreign_key_check").fetchall() == [])
            scan = db.execute("SELECT id FROM knowledge_documents WHERE title='消防集合扫描件'").fetchone()
            scan_text = "\n".join(row[0] for row in db.execute("SELECT text FROM knowledge_chunks WHERE document_id=? AND text_source='ocr'", (scan[0],)))
            images = db.execute("SELECT image_ref FROM knowledge_assets WHERE document_id=?", (scan[0],)).fetchall()
            check("真实扫描PDF OCR文字及图片落盘", "东侧广场" in "".join(scan_text.split()) and bool(images) and all((root / row[0]).is_file() for row in images), {"text": scan_text, "images": len(images)})
            service = app.state.assistant
            for role in ROLE_PERMISSIONS:
                user = next(a for a in accounts if a["role"] == role)
                hits = search_knowledge(db, user, "薪酬复核 财务专用标识 VPN运维", service.get_embedder(), service.get_store(), limit=30)
                document_scopes = dict(db.execute("SELECT id,scope FROM knowledge_documents"))
                check(f"{role}真实向量检索权限", all(f"kb:{document_scopes[h['document_id']]}" in ROLE_PERMISSIONS[role] for h in hits), {"hits": len(hits)})
            cases = [("如何申请星河公司的年假？", "星河年假制度"), ("薪酬调整需要谁审批？", "星河薪酬审核"),
                     ("VPN故障怎么办？", "星河VPN维护"), ("报销需要哪些材料？", "星河财务报销"),
                     ("消防集合地点在哪里？", "消防集合扫描件")]
            reranker = service.get_reranker()
            for question, expected in cases:
                hits = search_knowledge(db, actor, question, service.get_embedder(), service.get_store(), limit=12)
                if reranker and hits:
                    scores = reranker.scores(question, [h["text"] for h in hits])
                    hits = [h for _, h in sorted(zip(scores, hits), key=lambda pair: pair[0], reverse=True)]
                check(f"真实检索重排：{question}", bool(hits) and hits[0]["title"] == expected,
                      {"expected": expected, "actual": hits[0]["title"] if hits else None})
        if live:
            with closing(client_for()) as client:
                login(client, employee)
                cid = conversation(client)
                for question, expected_domain, marker in [
                    ("星河公司的年假申请流程是什么？", "knowledge", "主管"),
                    ("那需要提前多久申请？", "knowledge", "三"),
                    ("星河公司火星驻地补贴是多少？", "knowledge", "回复“转人工”"),
                    ("星河公司的财务专用标识是什么？", "knowledge", "回复“转人工”"),
                    ("我的电脑无法联网", "it", "确认报修"),
                    ("确认报修", "it", "已创建"),
                ]:
                    start = time.perf_counter()
                    response = send(client, cid, question)
                    payload = response.json()
                    # 知识缺少依据时先说明；低置信度允许请求澄清。
                    valid_domain = payload.get("domain") == expected_domain or (expected_domain == "knowledge" and payload.get("domain") == "clarify" and not payload.get("feedback_allowed"))
                    check(f"真实DeepSeek：{question}", response.status_code == 200 and valid_domain and marker in payload.get("answer", ""),
                          {"seconds": round(time.perf_counter() - start, 2), "http": response.status_code, "response": payload})
                    if payload.get("feedback_allowed"):
                        feedback_url = f"/api/messages/{payload['id']}/feedback"
                        check("真实知识回答反馈", client.post(feedback_url, json={"rating": "有用"}).status_code == 200)
                login(client, admin)
                response = send(client, conversation(client), "星河公司的财务专用标识是什么？")
                check("管理员真实模型可查财务专用资料", response.status_code == 200 and "FIN-LOTUS-964" in response.json().get("answer", ""), response.json())
                cookie = dict(client.cookies)
                parallel_ids = [conversation(client) for _ in range(5)]
                def parallel_knowledge(cid: str) -> dict[str, Any]:
                    started = time.perf_counter()
                    with closing(client_for()) as other:
                        other.cookies.update(cookie)
                        response = send(other, cid, "星河公司的年假申请流程是什么？")
                        return {"http": response.status_code, "seconds": round(time.perf_counter() - started, 2), "response": response.json()}
                started = time.perf_counter()
                with ThreadPoolExecutor(max_workers=5) as pool:
                    responses = list(pool.map(parallel_knowledge, parallel_ids))
                check("5并发真实知识问答", all(r["http"] == 200 and "主管" in r["response"].get("answer", "") and r["response"].get("feedback_allowed") for r in responses),
                      {"elapsed_seconds": round(time.perf_counter() - started, 2), "responses": responses})
        report = {"time": datetime.now().isoformat(), "transport": "HTTP" if base_url else "ASGI TestClient",
                  "concurrency": 10, "accounts": 100, "live_deepseek": live,
                  "counts": counts(root), "passed": sum(r["passed"] for r in results), "total": len(results), "results": results}
        save_json(root / f"report_{label}.json", report)
        (root / f"report_{label}.md").write_text(
            f"# 百人公司测试\n\n通过 {report['passed']}/{report['total']} 项；接口传输：{report['transport']}；10 并发。\n"
            "\n原始样本和预置对话是合成数据；只有标为真实 DeepSeek 的项目调用了模型。\n\n"
            + "\n".join(f"- {'通过' if r['passed'] else '**失败**'}：{r['name']}" for r in results)
            + "\n\n详见同名 JSON 中的响应证据、时间与统计。\n", encoding="utf-8")
        print(f"完成 {report['passed']}/{report['total']}：{root / f'report_{label}.json'}", flush=True)
        if report["passed"] != report["total"]:
            raise SystemExit(1)


def _other_conversation(root: Path, username: str) -> str:
    with sqlite3.connect(root / "company.db") as db:
        return db.execute("SELECT id FROM conversations WHERE username<>? LIMIT 1", (username,)).fetchone()[0]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("seed", "test", "serve"))
    parser.add_argument("--root", type=Path)
    parser.add_argument("--live", action="store_true", help="使用 .env 中真实 DeepSeek API，会消耗额度")
    parser.add_argument("--base-url", help="测试已启动的模拟环境真实 HTTP 地址")
    parser.add_argument("--port", type=int, default=8002)
    args = parser.parse_args()
    if args.command == "seed":
        seed()
        return
    if args.root is None:
        parser.error("test/serve 必须指定 --root")
    root = args.root.resolve()
    if args.command == "serve":
        uvicorn.run(simulation_app(root, args.port), host="127.0.0.1", port=args.port, log_level="info")
    else:
        run_tests(root, args.live, args.base_url)


if __name__ == "__main__":
    main()
