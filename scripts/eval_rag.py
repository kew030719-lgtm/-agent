"""用临时合成文件评估本地检索与重排，不调用 DeepSeek。"""

import os
import sqlite3
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf
from docx import Document
from openpyxl import Workbook

from app.auth import DB_PATH, User
from app.knowledge.answer import MIN_KNOWLEDGE_SCORE
from app.knowledge.documents import add_knowledge_file
from app.knowledge.embeddings import LocalEmbedder, index_document
from app.knowledge.reranker import LocalReranker
from app.knowledge.retrieval import SearchHit, search_knowledge
from app.knowledge.vector_store import MilvusVectorStore


def create_sample_files(root: Path) -> list[tuple[Path, str]]:
    markdown = root / "年假.md"
    markdown.write_text(
        "# 年假申请\n员工申请年假需要提交请假单，并由直属主管审批。",
        encoding="utf-8",
    )

    word_path = root / "VPN.docx"
    word = Document()
    word.add_heading("VPN 使用", 0)
    word.add_paragraph("电脑无法连接公司 VPN 时，请提交 IT 工单。")
    word.save(word_path)

    excel_path = root / "会议室.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "会议室"
    sheet.append(["名称", "容量", "位置"])
    sheet.append(["海棠会议室", 8, "三楼"])
    workbook.save(excel_path)

    pdf_path = root / "报销.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72), "报销申请需要附上发票并由财务审核。",
        fontname="china-s", fontsize=12,
    )
    pdf.save(pdf_path)
    pdf.close()

    return [
        (markdown, "general"), (word_path, "it"),
        (excel_path, "general"), (pdf_path, "finance"),
    ]


def rerank_hits(
    question: str,
    hits: list[SearchHit],
    reranker: LocalReranker,
) -> list[SearchHit]:
    scores = reranker.scores(question, [hit["text"] for hit in hits])
    return [
        hit for _score, hit in sorted(
            zip(scores, hits), key=lambda item: item[0], reverse=True
        )
    ]


def main() -> None:
    embedder = LocalEmbedder(DB_PATH.parent)
    reranker = LocalReranker(DB_PATH.parent)
    previous_uri = os.environ.get("MILVUS_URI")
    try:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            os.environ["MILVUS_URI"] = str(root / "vectors.db")
            with sqlite3.connect(":memory:") as db:
                db.execute("PRAGMA foreign_keys = ON")
                db.execute("CREATE TABLE users(username TEXT PRIMARY KEY)")
                db.execute("INSERT INTO users VALUES ('admin')")
                admin: User = {"username": "admin", "role": "admin"}
                employee: User = {"username": "employee", "role": "employee"}
                with MilvusVectorStore() as store:
                    for path, scope in create_sample_files(root):
                        document_id = add_knowledge_file(
                            db, admin, path, scope, root
                        )
                        index_document(db, document_id, embedder, store)

                    cases = [
                        ("如何申请年假", "年假"),
                        ("电脑 VPN 无法连接怎么办", "VPN"),
                        ("海棠会议室可以坐几个人", "会议室"),
                        ("报销申请需要什么材料", "报销"),
                    ]
                    top_one = 0
                    eligible = 0
                    for question, expected in cases:
                        hits = search_knowledge(
                            db, admin, question, embedder, store, limit=4
                        )
                        ranked = rerank_hits(question, hits, reranker)
                        actual = ranked[0]["title"] if ranked else "无结果"
                        top_one += actual == expected
                        score = ranked[0]["score"] if ranked else 0.0
                        eligible += actual == expected and score >= MIN_KNOWLEDGE_SCORE
                        print(f"{question} -> {actual}（期望：{expected}，余弦分数：{score:.3f}）")

                    for question in ("公司婚假一共有几天？", "公司停车补贴每月多少钱？"):
                        candidates = search_knowledge(db, admin, question, embedder, store, limit=4)
                        passed = sum(hit["score"] >= MIN_KNOWLEDGE_SCORE for hit in candidates)
                        print(f"无答案问题：{question}，达到旧0.55诊断线的候选：{passed}（不是当前问答的硬门槛，仍需生成阶段判断证据）")

                    restricted = search_knowledge(
                        db, employee, "报销申请需要什么材料",
                        embedder, store, limit=4,
                    )
                    leaked = any(
                        hit["title"] == "报销" for hit in restricted
                    )
                    print(f"top-1：{top_one}/{len(cases)}，达到旧分数诊断线：{eligible}/{len(cases)}，权限泄漏：{int(leaked)}")
                    if top_one != len(cases) or leaked:
                        raise RuntimeError("合成数据检索评估未通过")
    finally:
        if previous_uri is None:
            os.environ.pop("MILVUS_URI", None)
        else:
            os.environ["MILVUS_URI"] = previous_uri


if __name__ == "__main__":
    main()
